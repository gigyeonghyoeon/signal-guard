# -*- coding: utf-8 -*-
"""Convert Markdown requirements doc to Word .docx"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Cm, RGBColor


def set_run_font(run, size=10, bold=False, code=False):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Consolas" if code else "맑은 고딕"
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    if code:
        rFonts.set(qn("w:ascii"), "Consolas")
        rFonts.set(qn("w:hAnsi"), "Consolas")
        rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    else:
        rFonts.set(qn("w:ascii"), "맑은 고딕")
        rFonts.set(qn("w:hAnsi"), "맑은 고딕")
        rFonts.set(qn("w:eastAsia"), "맑은 고딕")


def add_inline(paragraph, text: str, base_size=10):
    """Parse **bold**, `code`, and [text](url) into runs."""
    pattern = re.compile(
        r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))"
    )
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos : m.start()])
            set_run_font(run, size=base_size)
        token = m.group(0)
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, bold=True)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=base_size - 0.5, code=True)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        elif token.startswith("["):
            link_m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link_m:
                label, url = link_m.group(1), link_m.group(2)
                run = paragraph.add_run(f"{label} ({url})")
                set_run_font(run, size=base_size)
                run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
                run.underline = True
            else:
                run = paragraph.add_run(token)
                set_run_font(run, size=base_size)
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=base_size)


def set_cell_shading(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.autofit = True
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.rows[i].cells[j]
            text = row[j] if j < len(row) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, text.strip(), base_size=9)
            for para in cell.paragraphs:
                for run in para.runs:
                    set_run_font(run, size=9, bold=(i == 0 and run.bold))
            if i == 0:
                set_cell_shading(cell, "D9E2F3")
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph()


def parse_table_block(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        line = lines[i].strip()
        # separator row
        if re.match(r"^\|[\s\-:|]+\|$", line):
            i += 1
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        rows.append(cells)
        i += 1
    return rows, i


def convert(md_path: Path, docx_path: Path):
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()

    # page margins
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    # default style
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(10)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15

    heading_sizes = {1: 18, 2: 14, 3: 12, 4: 11, 5: 10.5, 6: 10}

    i = 0
    in_code = False
    code_lines: list[str] = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # code fence
        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                in_code = False
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.left_indent = Cm(0.3)
                run = p.add_run("\n".join(code_lines))
                set_run_font(run, size=8.5, code=True)
                # light background via shading on paragraph
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "F2F2F2")
                shd.set(qn("w:val"), "clear")
                p._p.get_or_add_pPr().append(shd)
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        # blank
        if not stripped:
            i += 1
            continue

        # horizontal rule
        if re.match(r"^-{3,}$", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "A6A6A6")
            pBdr.append(bottom)
            pPr.append(pBdr)
            i += 1
            continue

        # heading
        hm = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if hm:
            level = len(hm.group(1))
            title = hm.group(2).strip()
            p = doc.add_heading(level=min(level, 9))
            p.clear()
            run = p.add_run(title)
            set_run_font(run, size=heading_sizes[level], bold=True)
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # table
        if stripped.startswith("|"):
            rows, next_i = parse_table_block(lines, i)
            add_table(doc, rows)
            i = next_i
            continue

        # checklist
        cm = re.match(r"^- \[([ xX])\]\s+(.*)$", stripped)
        if cm:
            mark = "☑" if cm.group(1).lower() == "x" else "☐"
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            add_inline(p, f"{mark} {cm.group(2)}")
            i += 1
            continue

        # unordered list
        um = re.match(r"^[-*]\s+(.*)$", stripped)
        if um:
            p = doc.add_paragraph(style="List Bullet")
            p.clear()
            add_inline(p, um.group(1))
            i += 1
            continue

        # ordered list
        om = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if om:
            p = doc.add_paragraph(style="List Number")
            p.clear()
            add_inline(p, om.group(2))
            i += 1
            continue

        # normal paragraph
        p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1

    doc.save(docx_path)
    print(f"Saved: {docx_path}")


if __name__ == "__main__":
    import sys

    base = Path(__file__).resolve().parent
    default_files = [
        "유스케이스사용자스토리.md",
        "제안요청서요구사항대응표.md",
        "백로그WBS.md",
    ]
    targets = sys.argv[1:] or default_files
    for name in targets:
        md = base / name
        convert(md, base / f"{md.stem}.docx")
